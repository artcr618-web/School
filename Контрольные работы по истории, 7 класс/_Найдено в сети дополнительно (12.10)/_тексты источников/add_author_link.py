# -*- coding: utf-8 -*-
"""12.10.2026: убрать столбец «Примечание», добавить столбцы «Автор» и «Ссылка на источник».
Перед удалением содержимое «Примечаний» сохраняется отдельным файлом.
"""
import copy, urllib.parse
from openpyxl import load_workbook
from openpyxl.styles import Font

DIR = "/home/user/School/Контрольные тестовые работы/"
PATH = DIR + "Контрольные работы 7 класс.xlsx"
NOTES_OUT = DIR + "_Примечания (убраны из таблицы 12.10.2026).md"

S_BASE = "https://school105.ru/storage_drive/svedeniya-ob-oo/osnovnye-svedeniya/obrazovanie/om/5-9/История/"
L_NS_GALSANOV = "https://nsportal.ru/vuz/istoricheskie-nauki/library/2026/04/09/testy-po-istorii-rossii-7-klass"
L_NS_NIKITINA = "https://nsportal.ru/shkola/istoriya/library/2019/10/09/test-po-istorii-rossii"
A_GALSANOV = "Галсанов Ринчин-Доржи Пурбоевич"
A_NIKITINA = "Никитина Людмила Александровна"
A_NONE = "автор не указан (school105.ru)"

def s105(fname):
    return S_BASE + fname

LINK_MAP = [
    ("Итоговый тест «История России", A_GALSANOV, L_NS_GALSANOV),
    ("Тест «Россия в XVI в.»", A_GALSANOV, L_NS_GALSANOV),
    ("Тест «Смута в России»", A_GALSANOV, L_NS_GALSANOV),
    ("Тест «Россия при первых Романовых»", A_GALSANOV, L_NS_GALSANOV),
    ("Тест «Россия в XVII веке»", A_NONE, s105("Контрольная работа по истории России 7 класс.pdf")),
    ("Итоговая контрольная работа по Новой истории", A_NONE, s105("Контрольная работа по истории 7 класс новая история.pdf")),
    ("Тест «Россия на рубеже XVI–XVII веков»", A_NIKITINA, L_NS_NIKITINA),
    ("Контрольная работа по истории России XVI–XVII вв.", A_NONE, s105("ОМ история 7 класс.pdf")),
    ("Итоговая работа по всеобщей истории", A_NONE, s105("ОМ история 7 класс.pdf")),
    ("Входная контрольная работа, 7 класс", A_NONE, s105("Входная контрольная работа по истории 7 класс.pdf")),
    ("Входная контрольная работа, 8 класс", A_NONE, s105("Входная контрольная работа по истории 8 класс.pdf")),
]

wb = load_workbook(PATH)
ws = wb["Работы"]

# 1) сохранить примечания
notes = []
for r in ws.iter_rows(min_row=2, values_only=True):
    if r[13]:
        notes.append(r)
with open(NOTES_OUT, "w", encoding="utf-8") as f:
    f.write("# Пояснения, которые были в столбце «Примечание» (убраны 12.10.2026)\n\n")
    f.write("Столбец «Примечание» убран из таблицы по просьбе заказчика; текст сохранён здесь.\n")
    f.write("Здесь — расшифровки ответов составных заданий, откуда взята работа и пометки\nо расхождениях с ключами источников.\n\n")
    f.write("| № в таблице | Работа | Вариант | № задания | Примечание |\n|---|---|---|---|---|\n")
    for r in notes:
        txt = str(r[13]).replace("|", "/")
        f.write("| %s | %s | %s | %s | %s |\n" % (r[0], r[2], r[3], r[4], txt))
print("сохранено примечаний:", len(notes))

# 2) заменить шапку
for col, title in ((14, "Автор"), (15, "Ссылка на источник")):
    ref = ws.cell(row=1, column=1)
    c = ws.cell(row=1, column=col)
    c.font = copy.copy(ref.font); c.fill = copy.copy(ref.fill)
    c.border = copy.copy(ref.border); c.alignment = copy.copy(ref.alignment)
    c.value = title

# 3) заполнить строки
unmatched = set()
for i, r in enumerate(ws.iter_rows(min_row=2), start=2):
    work = str(r[2].value)
    author = link = None
    for key, a, l in LINK_MAP:
        if work.startswith(key) or key in work:
            author, link = a, l
            break
    if author is None:
        unmatched.add(work)
        author, link = "", ""
    for col, val in ((14, author), (15, link)):
        ref = ws.cell(row=2, column=col)
        c = ws.cell(row=i, column=col)
        c.font = copy.copy(ref.font); c.border = copy.copy(ref.border)
        c.fill = copy.copy(ref.fill); c.alignment = copy.copy(ref.alignment)
        c.number_format = ref.number_format
        c.value = val
    c = ws.cell(row=i, column=15)
    if link:
        c.hyperlink = urllib.parse.quote(link, safe=":/?&=%")
        c.font = Font(color="0563C1", underline="single")

ws.column_dimensions["N"].width = 30
ws.column_dimensions["O"].width = 62
ws.auto_filter.ref = "A1:O%d" % ws.max_row
wb.save(PATH)
print("строк с автором/ссылкой:", ws.max_row - 1, "| не найдено соответствий:", unmatched)
