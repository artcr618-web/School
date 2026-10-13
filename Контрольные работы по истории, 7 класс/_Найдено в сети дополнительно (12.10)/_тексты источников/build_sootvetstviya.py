# -*- coding: utf-8 -*-
"""Собирает лист «Соответствия» (37 заданий) и убирает их из листа «Работы»."""
import json
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

DIR = "/home/user/School/Контрольные тестовые работы/"
PATH = DIR + "Контрольные работы 7 класс.xlsx"
T = json.load(open(DIR + "_Найдено в сети дополнительно (12.10)/_тексты источников/sootvetstviya.json"))

def clean_body(s):
    s = (s or "").strip()
    s = s.strip("|").strip()
    s = s.rstrip(";").strip()
    if s.endswith("|"): s = s.rstrip("|").strip()
    return s

def cell_label(label, body):
    return "%s) %s" % (label, clean_body(body))

def answer_text(t):
    order = [l for l, _ in t["left"]]
    parts = []
    for l in order:
        v = t["pairs"].get(l)
        if v is None: continue
        if len(v) > 1 and all(ch in [x for x, _ in t["right"]] for ch in v):
            v = ", ".join(v)
        parts.append("%s → %s" % (l, v))
    return "; ".join(parts)

wb = load_workbook(PATH)
ws = wb["Работы"]

nums = [t["num"] for t in T]
# 1) удаляем строки соответствий из «Работы» (снизу вверх)
for t in sorted(T, key=lambda x: -x["num"]):
    ws.delete_rows(int(t["num"]) + 1)
# 2) перенумеровываем
for i, r in enumerate(ws.iter_rows(min_row=2), start=1):
    r[0].value = i
ws.auto_filter.ref = "A1:R%d" % ws.max_row
print("в «Работах» осталось:", ws.max_row - 1)

# 3) новый лист
if "Соответствия" in wb.sheetnames:
    del wb["Соответствия"]
sh = wb.create_sheet("Соответствия", 1)

HEAD = ["№", "Раздел", "Работа", "Вариант", "№ задания", "Вопрос",
        "Левая часть 1", "Левая часть 2", "Левая часть 3", "Левая часть 4", "Левая часть 5",
        "Правая часть 1", "Правая часть 2", "Правая часть 3", "Правая часть 4", "Правая часть 5",
        "Правильный ответ", "Ответ как в источнике", "Баллы", "Примечание"]
WID = [5, 14, 34, 8, 9, 46, 28, 28, 28, 28, 28, 28, 28, 28, 28, 28, 30, 22, 7, 32]
hfont = Font(name="Calibri", size=11, bold=True)
hfill = PatternFill("solid", fgColor="DDEBF7")
dfont = Font(name="Calibri", size=11)
dalign = Alignment(wrap_text=True, vertical="top")

for j, h in enumerate(HEAD, start=1):
    c = sh.cell(row=1, column=j, value=h)
    c.font = hfont; c.fill = hfill; c.alignment = Alignment(vertical="top", wrap_text=True)
    sh.column_dimensions[get_column_letter(j)].width = WID[j-1]

for i, t in enumerate(sorted(T, key=lambda x: x["num"]), start=2):
    row = [i-1, t["razdel"], t["work"], t["var"], t["task"], t["head"]]
    for n in range(5):
        row.append(cell_label(*t["left"][n]) if n < len(t["left"]) else None)
    for n in range(5):
        row.append(cell_label(*t["right"][n]) if n < len(t["right"]) else None)
    row += [answer_text(t), t["answer"], t["ball"], t["note"]]
    for j, v in enumerate(row, start=1):
        c = sh.cell(row=i, column=j, value=v)
        c.font = dfont; c.alignment = dalign

sh.freeze_panes = "G2"
sh.auto_filter.ref = "A1:T%d" % sh.max_row
if "Соответствия" in wb.sheetnames and wb.sheetnames[0] != "Работы":
    wb.move_sheet("Работы", offset=-10)
wb.save(PATH)
print("лист «Соответствия»: строк", sh.max_row - 1, "| порядок листов:", wb.sheetnames)
