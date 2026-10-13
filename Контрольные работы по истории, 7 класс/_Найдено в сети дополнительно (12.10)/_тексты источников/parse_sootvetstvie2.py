# -*- coding: utf-8 -*-
"""Разбор 37 заданий «на соответствие» — версия 2 (учёт «г.» в датах, инициалов «Б. Годунове»,
разных стилей маркеров). Результат — sootvetstviya.json + печать для глазной проверки."""
import re, json
from openpyxl import load_workbook

DIR = "/home/user/School/Контрольные тестовые работы/"
wb = load_workbook(DIR + "Контрольные работы 7 класс.xlsx")
ws = wb["Работы"]
ROWS = list(ws.iter_rows(min_row=2, values_only=True))

LAT2CYR = {"A": "А", "B": "Б", "C": "В", "D": "Г", "E": "Д"}
CYR = "АБВГДЕ"
INTRO = ["Дата", "Даты", "Определение", "Определения", "Определение понятия", "Деятельность",
         "Роль", "Участник", "Имя", "Имена", "Событие", "События", "Название", "Личность",
         "Термин", "Понятие", "Значение", "Результат", "Война", "Войны", "Характеристика",
         "Исторический деятель", "Экономические и политические особенности земли"]

def count_style(text, kind):
    if kind == "up_paren":
        return len(re.findall(r'(?<![А-Яа-яЁёA-Za-z0-9])(?:[А-ЯA-Z])\s*\)\s', text))
    if kind == "up_dot":
        return len(re.findall(r'(?<![А-Яа-яЁёA-Za-z0-9])(?:[А-ЯA-Z])\s*\.\s(?=[А-ЯЁA-Z«"])', text))
    if kind == "low_paren":
        return len(re.findall(r'(?<![А-Яа-яЁёA-Za-z])(?:[абвгде])\s*\)\s', text))
    if kind == "dig_paren":
        return len(re.findall(r'(?<!\d)(\d{1,2})\s*\)\s', text))
    if kind == "dig_dot":
        return len(re.findall(r'(?<![.\d-])(\d{1,2})\s*\.\s(?=[А-ЯЁA-Z])', text))
    return 0

RX = {
  "up_paren":  re.compile(r'(?<![А-Яа-яЁёA-Za-z0-9])(?:[А-ЯA-Z])\s*\)\s'),
  "up_dot":    re.compile(r'(?<![А-Яа-яЁёA-Za-z0-9])(?:[А-ЯA-Z])\s*\.\s(?=[А-ЯЁA-Z«"])'),
  "low_paren": re.compile(r'(?<![А-Яа-яЁёA-Za-z])(?:[абвгде])\s*\)\s'),
  "dig_paren": re.compile(r'(?<!\d)(\d{1,2})\s*\)\s'),
  "dig_dot":   re.compile(r'(?<![.\d-])(\d{1,2})\s*\.\s(?=[А-ЯЁA-Z])'),
}

def pick_styles(q):
    """выбирает пару видов маркеров (левая и правая части)"""
    c = {k: count_style(q, k) for k in RX}
    def first_pos(key):
        m = RX[key].search(q)
        return m.start() if m else 10**9
    if c["dig_paren"] >= 2 and c["low_paren"] >= 3:
        pair = ["dig_paren", "low_paren"]
    elif c["up_paren"] >= 3 and c["dig_paren"] >= 2:
        pair = ["up_paren", "dig_paren"]
    elif c["up_dot"] >= 3 and c["dig_paren"] >= 2:
        pair = ["up_dot", "dig_paren"]
    elif c["dig_paren"] >= 2 and c["up_dot"] >= 3:
        pair = ["dig_paren", "up_dot"]
    elif c["dig_dot"] >= 3 and c["up_paren"] >= 3:
        pair = ["dig_dot", "up_paren"]
    elif c["dig_paren"] >= 2 and c["up_paren"] >= 3:
        pair = ["dig_paren", "up_paren"]
    else:
        pair = sorted([k for k in ("up_paren", "dig_paren", "low_paren") if c[k] >= 2],
                      key=first_pos)[:2]
    pair = sorted(pair, key=first_pos)   # кто первый в тексте — тот слева
    return pair, c

def norm_label(raw):
    ch = raw.strip()[0]
    if ch in LAT2CYR: ch = LAT2CYR[ch]
    return ch

def tokenize(text, pair):
    rxs = [RX[k] for k in pair]
    ms = []
    for k, rx in zip(pair, rxs):
        for m in rx.finditer(text):
            ms.append((m.start(), m.end(), k, m.group(0)))
    ms.sort()
    parts = []
    for i, (st, en, k, raw) in enumerate(ms):
        end = ms[i+1][0] if i+1 < len(ms) else len(text)
        label = norm_label(raw)
        body = text[en:end].strip()
        body = re.sub(r'^\|+', '', body).strip()
        body = re.sub(r'\|+$', '', body).strip()
        parts.append((k, label, body))
    head = text[:ms[0][0]].strip() if ms else text
    return head, parts

def trim_intro(s):
    s = s.strip()
    changed = True
    while changed:
        changed = False
        for w in sorted(INTRO, key=len, reverse=True):
            if re.search(r'(?:^|\s)' + re.escape(w) + r'\s*[;.:]?\s*$', s):
                s = re.sub(r'(?:^|\s)' + re.escape(w) + r'\s*[;.:]?\s*$', '', s).strip()
                changed = True
    return s

def is_letter(l): return l in CYR or l in "абвгде"


def parse_task(num, q):
    pair, counts = pick_styles(q)
    head, parts = tokenize(q, pair)
    parts = [(k, l, b) for k, l, b in parts if b.strip()]
    if not parts: return None, None
    kind_left = parts[0][0]
    left, right = [], []
    for k, label, body in parts:
        (left if k == kind_left else right).append([label, body])
    if left: left[-1][1] = trim_intro(left[-1][1])
    head = trim_intro(head)
    return (head, left, right), dict(pair=pair, counts=counts)

def parse_answer(a, left, right):
    a = (a or "").strip()
    lab_r = [l for l, _ in right]
    def fix(letter):
        for L in lab_r:
            if L.lower() == letter.lower(): return L
        return letter
    pairs = {}
    m = re.findall(r'([А-ЯA-Zа-я]|\d)\s*[—\-–]+\s*([А-ЯA-Zа-я]|\d)', a)
    if len(m) >= 2:
        for l, r in m:
            L = l.upper() if l.isalpha() and l.isupper() else l
            pairs[l] = fix(r) if r.isalpha() else r
        return pairs, "через тире"
    m = re.findall(r'([А-ЯA-Z])\s*(\d)', a)
    if len(m) >= 2:
        for l, r in m: pairs[l] = r
        return pairs, "буква+цифра"
    if re.fullmatch(r'\d{2,6}', a):
        labs = [l for l, _ in left]
        if len(a) == len(labs):
            for l, r in zip(labs, a): pairs[l] = r
            return pairs, "строка цифр"
    m = re.findall(r'(\d)\s*([А-ЯA-Zа-я]{1,4})(?![а-я])', a)
    if m:
        for l, r in m: pairs[l] = r
        return pairs, "цифра+буквы"
    m = re.findall(r'(\d)\s*[\)\.]?\s*([абвгде])', a)
    if m:
        for l, r in m: pairs[l] = fix(r)
        return pairs, "цифра+буква"
    return {}, "НЕ РАСПОЗНАН"

# ручные правки для заданий с опечатками/особым форматом в источнике
MANUAL = {
 241: dict(  # uvarova вар.2 В3: в источнике «2» без скобки; собран вручную
   head="Установите соответствие между понятием и его определением.",
   left=[["1", "парсуна"],
         ["2", "новое направление в архитектуре XVII века"],
         ["3", "сатирическая повесть"],
         ["4", "житие"]],
   right=[["а", "произведение, в котором обличались феодальные порядки, крючкотворство, продажность чиновников, волокита"],
          ["б", "«нарышкинское барокко»"],
          ["в", "светский портрет"],
          ["г", "биографическая повесть"]],
   pairs={"1": "в", "2": "б", "3": "а", "4": "г"}, how="через тире (правка вручную)"),
}

out = []
for r in ROWS:
    q = str(r[5] or "")
    if not (all(o is None for o in r[6:12]) and re.search(r'соответстви|Соотнесите', q)):
        continue
    num = r[0]
    if num in MANUAL:
        m = MANUAL[num]
        head, left, right, pairs, how, styles = m["head"], m["left"], m["right"], m["pairs"], m["how"], {}
    else:
        p, styles = parse_task(num, q)
        if not p:
            print('!! не разобрано:', num); continue
        head, left, right = p
        pairs, how = parse_answer(str(r[12] or ""), left, right)
    out.append(dict(num=num, razdel=r[1], work=r[2], var=r[3], task=r[4], head=head,
                    left=left, right=right, answer=str(r[12] or ""), pairs=pairs, how=how,
                    styles=styles, ball=r[13], note=r[17]))

print("РАЗОБРАНО:", len(out))
bad = []
for t in out:
    src = t["answer"]
    ok = len(t["pairs"]) and all(str(k) in [l for l, _ in t["left"]] for k in t["pairs"])
    if not ok: bad.append(t["num"])
    print("-"*100)
    print("№%s | %s | вар.%s | зад.%s | стили: %s | ответ распознан как: %s %s" %
          (t["num"], str(t["work"])[:38], t["var"], t["task"], t["styles"], t["how"], "ОК" if ok else "!!! ПРОВЕРИТЬ"))
    print("   ВОПРОС:", t["head"][:110])
    for l, b in t["left"]:  print("   Л %s: %s" % (l, b[:95]))
    for l, b in t["right"]: print("   П %s: %s" % (l, b[:95]))
    print("   ОТВЕТ ИСТОЧНИКА:", src, "=>", t["pairs"])
print("\nПРОБЛЕМНЫЕ:", bad)
json.dump(out, open(DIR + "_Найдено в сети дополнительно (12.10)/_тексты источников/sootvetstviya.json", "w"),
          ensure_ascii=False, indent=1, default=str)
